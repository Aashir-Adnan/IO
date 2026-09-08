from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from io_assistant.config import FilesConfig

__all__ = ["ExtractedDoc", "read_document"]

_HEAD_MD = re.compile(r"^#\s+(.+)$", re.M)
_HEAD_HTML = re.compile(r"<h1[^>]*>(.*?)</h1>", re.I | re.S)
_TAGS = re.compile(r"<[^>]+>")
_RTF = re.compile(r"\\[a-z]+-?\d* ?|[{}]")


@dataclass
class ExtractedDoc:
    text: str
    title: str | None
    truncated: bool
    status: str
    error: str | None = None


def _title_from(text: str, suffix: str) -> str | None:
    if suffix in (".md", ".markdown", ".rst"):
        m = _HEAD_MD.search(text)
        if m:
            return m.group(1).strip()
    if suffix in (".html", ".htm"):
        m = _HEAD_HTML.search(text)
        if m:
            return _TAGS.sub("", m.group(1)).strip()
    return None


def read_document(path: Path, config: FilesConfig) -> ExtractedDoc:
    try:
        size = path.stat().st_size
    except OSError as e:
        return ExtractedDoc("", None, False, "unreadable", str(e))
    if size > config.max_file_bytes:
        return ExtractedDoc("", None, False, "too_large", None)

    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            text = "\n".join((pg.extract_text() or "") for pg in reader.pages)
            title = (reader.metadata.title if reader.metadata else None) or None
        elif suffix == ".docx":
            import docx

            doc = docx.Document(str(path))
            text = "\n".join(p.text for p in doc.paragraphs)
            title = doc.paragraphs[0].text.strip() if doc.paragraphs else None
        elif suffix in (".html", ".htm", ".xml"):
            raw = path.read_text(encoding="utf-8", errors="replace")
            title = _title_from(raw, suffix)
            text = re.sub(r"\s+\n", "\n", _TAGS.sub(" ", raw))
        elif suffix == ".rtf":
            raw = path.read_text(encoding="utf-8", errors="replace")
            text, title = _RTF.sub("", raw), None
        elif suffix in (".csv", ".tsv"):
            lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
            text, title = "\n".join(lines[:200]), None
        else:
            try:
                text = path.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                text = path.read_bytes().decode("latin-1")
            title = _title_from(text, suffix)
    except Exception as e:  # pragma: no cover - defensive
        return ExtractedDoc("", None, False, "unreadable", f"{type(e).__name__}: {e}")

    sample = text[:4096]
    if "\x00" in sample or (
        sample
        and sum(c.isprintable() or c in "\r\n\t" for c in sample) / len(sample) < 0.85
    ):
        return ExtractedDoc("", title, False, "unreadable", "binary content")

    text = text.strip()
    if not text:
        return ExtractedDoc("", title, False, "unreadable", "no extractable text")
    if title is None:
        title = path.stem
    return ExtractedDoc(text, title, False, "ok", None)
